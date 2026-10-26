from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.dependencies import get_rate_limiter, get_store
from app.main import app
from app.services.excel_store import COLUMNS, SHEET_NAME, ExcelEnquiryStore
from app.services.rate_limiter import RateLimiter

VALID = {
    "name": "Jane Cooper",
    "email": "jane@company.com",
    "project_types": ["website", "ai"],
    "country": "GB",
    "budget_amount": 25000,
    "budget_not_sure": False,
    "timeline": "1_3_months",
    "message": "We want an AI assistant for our support team.",
    "partner_program": True,
    "consent": True,
    "website": "",
}

# Column positions in the spreadsheet (see COLUMNS in excel_store.py).
COL = {column.header: index for index, column in enumerate(COLUMNS)}


@pytest.fixture()
def excel_path(tmp_path: Path) -> Path:
    return tmp_path / "enquiries.xlsx"


@pytest.fixture()
def client(excel_path: Path):
    limiter = RateLimiter(max_requests=3, window_seconds=60)
    app.dependency_overrides[get_store] = lambda: ExcelEnquiryStore(excel_path)
    app.dependency_overrides[get_rate_limiter] = lambda: limiter
    yield TestClient(app)
    app.dependency_overrides.clear()


def read_rows(path: Path):
    sheet = load_workbook(path)[SHEET_NAME]
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def first_entry(path: Path) -> dict:
    """The first data row as {header: value}."""
    return dict(zip(COL, read_rows(path)[1]))


def test_health(client: TestClient) -> None:
    assert client.get("/api/health").json() == {"status": "ok"}


def test_valid_enquiry_is_written_to_excel(client: TestClient, excel_path: Path) -> None:
    response = client.post("/api/enquiries", json=VALID)

    assert response.status_code == 201
    reference = response.json()["id"]
    assert reference.startswith("TH-")

    rows = read_rows(excel_path)
    assert rows[0] == [column.header for column in COLUMNS]
    assert rows[1][1:] == [
        reference,
        "New",
        "Jane Cooper",
        "jane@company.com",
        "United Kingdom",
        "Website, AI integration",
        "GBP",
        25000,
        "1–3 months",
        "Yes",
        "We want an AI assistant for our support team.",
    ]


def test_budget_is_a_real_number_in_excel(client: TestClient, excel_path: Path) -> None:
    client.post("/api/enquiries", json=VALID)
    cell = load_workbook(excel_path)[SHEET_NAME].cell(row=2, column=COL["Budget"] + 1)
    assert cell.data_type == "n"
    assert cell.number_format == "#,##0"


def test_rows_are_appended_not_overwritten(client: TestClient, excel_path: Path) -> None:
    client.post("/api/enquiries", json=VALID)
    client.post("/api/enquiries", json={**VALID, "name": "Second Person"})
    rows = read_rows(excel_path)
    assert len(rows) == 3
    assert rows[2][COL["Name"]] == "Second Person"


def test_formula_like_text_is_stored_as_plain_text(client: TestClient, excel_path: Path) -> None:
    client.post("/api/enquiries", json={**VALID, "message": "=HYPERLINK(\"http://evil\")"})
    sheet = load_workbook(excel_path)[SHEET_NAME]
    cell = sheet.cell(row=2, column=len(COLUMNS))
    assert cell.data_type == "s"
    assert cell.value == "=HYPERLINK(\"http://evil\")"


@pytest.mark.parametrize(
    "country, amount, expected_country, expected_currency",
    [
        ("US", 1000, "United States", "USD"),
        ("GB", 4000, "United Kingdom", "GBP"),
        ("IN", 400000, "India", "INR"),
        ("AE", 18000, "United Arab Emirates", "AED"),
    ],
)
def test_budget_is_stored_with_local_currency(
    client: TestClient, excel_path: Path, country: str, amount: int, expected_country: str, expected_currency: str
) -> None:
    client.post("/api/enquiries", json={**VALID, "country": country, "budget_amount": amount})
    entry = first_entry(excel_path)
    assert (entry["Country"], entry["Currency"], entry["Budget"]) == (expected_country, expected_currency, amount)


def test_not_sure_budget(client: TestClient, excel_path: Path) -> None:
    response = client.post("/api/enquiries", json={**VALID, "budget_amount": None, "budget_not_sure": True})
    assert response.status_code == 201
    assert first_entry(excel_path)["Budget"] == "Not sure yet"


def test_amount_is_ignored_when_not_sure(client: TestClient, excel_path: Path) -> None:
    client.post("/api/enquiries", json={**VALID, "budget_amount": 5000, "budget_not_sure": True})
    assert first_entry(excel_path)["Budget"] == "Not sure yet"


def test_budget_is_required_unless_not_sure(client: TestClient) -> None:
    payload = {key: value for key, value in VALID.items() if key != "budget_amount"}
    response = client.post("/api/enquiries", json=payload)
    assert response.status_code == 422
    assert "budget_amount" in response.json()["detail"][0]["loc"]


def test_other_country_requires_a_name(client: TestClient, excel_path: Path) -> None:
    response = client.post("/api/enquiries", json={**VALID, "country": "OTHER"})
    assert response.status_code == 422
    assert "other_country" in response.json()["detail"][0]["loc"]

    client.post("/api/enquiries", json={**VALID, "country": "OTHER", "other_country": "Japan", "budget_amount": 8000})
    entry = first_entry(excel_path)
    assert (entry["Country"], entry["Currency"], entry["Budget"]) == ("Other: Japan", "USD", 8000)


def test_other_country_is_ignored_for_listed_countries(client: TestClient, excel_path: Path) -> None:
    client.post("/api/enquiries", json={**VALID, "other_country": "Japan"})
    assert first_entry(excel_path)["Country"] == "United Kingdom"


@pytest.mark.parametrize(
    "field, value",
    [
        ("name", "J"),
        ("email", "not-an-email"),
        ("project_types", []),
        ("project_types", ["spaceship"]),
        ("project_types", "website"),
        ("country", "XX"),
        ("country", ""),
        ("budget_amount", 0),
        ("budget_amount", -500),
        ("budget_amount", 12.5),
        ("budget_amount", "lots"),
        ("budget_amount", 10**12),
        ("timeline", "yesterday"),
        ("consent", False),
        ("message", "x" * 2001),
    ],
)
def test_invalid_fields_are_rejected(client: TestClient, excel_path: Path, field: str, value: object) -> None:
    response = client.post("/api/enquiries", json={**VALID, field: value})
    assert response.status_code == 422
    assert field in response.json()["detail"][0]["loc"]
    assert not excel_path.exists()


def test_unknown_fields_are_rejected(client: TestClient) -> None:
    assert client.post("/api/enquiries", json={**VALID, "admin": True}).status_code == 422


def test_duplicate_project_types_are_collapsed(client: TestClient, excel_path: Path) -> None:
    client.post("/api/enquiries", json={**VALID, "project_types": ["ai", "ai", "website"]})
    assert first_entry(excel_path)["Project types"] == "AI integration, Website"


def test_honeypot_submissions_are_not_stored(client: TestClient, excel_path: Path) -> None:
    response = client.post("/api/enquiries", json={**VALID, "website": "http://spam.example"})
    assert response.status_code == 201
    assert not excel_path.exists()


def test_rate_limit(client: TestClient) -> None:
    statuses = [client.post("/api/enquiries", json=VALID).status_code for _ in range(4)]
    assert statuses == [201, 201, 201, 429]


def test_rate_limiter_window_expires() -> None:
    now = [0.0]
    limiter = RateLimiter(max_requests=1, window_seconds=10, clock=lambda: now[0])
    assert limiter.allow("ip")
    assert not limiter.allow("ip")
    now[0] = 10.0
    assert limiter.allow("ip")
