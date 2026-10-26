"""Stores enquiries as rows in an Excel workbook.

Design notes:
- A file lock serialises writes, so several worker processes can safely share one file.
- Each save goes to a temporary file that then replaces the original, so a crash
  mid-write can never leave a half-written workbook behind.
- Every text value is forced to be a plain string cell. openpyxl would otherwise
  treat text that starts with "=" as a formula (spreadsheet formula injection).
"""

from __future__ import annotations

import logging
import os
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional, Sequence, Union

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

SHEET_NAME = "Enquiries"


@dataclass(frozen=True)
class Column:
    header: str
    width: int


# Order here defines the column order in the spreadsheet.
COLUMNS: Sequence[Column] = (
    Column("Submitted (UTC)", 20),
    Column("Reference", 14),
    Column("Status", 12),
    Column("Name", 24),
    Column("Email", 30),
    Column("Country", 22),
    Column("Project types", 28),
    Column("Currency", 10),
    Column("Budget", 14),
    Column("Timeline", 20),
    Column("Founding partner", 16),
    Column("Message", 60),
)


@dataclass(frozen=True)
class EnquiryRecord:
    reference: str
    submitted_at: datetime
    name: str
    email: str
    country: str
    project_type: str
    currency: str
    # Whole units of `currency`; None means the visitor ticked "Not sure yet".
    budget: Optional[int]
    timeline: str
    founding_partner: bool
    message: str

    def as_row(self) -> List[Union[str, int]]:
        return [
            self.submitted_at.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
            self.reference,
            "New",
            self.name,
            self.email,
            self.country,
            self.project_type,
            self.currency,
            self.budget if self.budget is not None else "Not sure yet",
            self.timeline,
            "Yes" if self.founding_partner else "No",
            self.message,
        ]


class ExcelStoreError(RuntimeError):
    """Raised when an enquiry could not be written to the workbook."""


class ExcelEnquiryStore:
    def __init__(self, path: Path, lock_timeout_seconds: float = 10.0) -> None:
        self.path = path
        self._lock = FileLock(str(path) + ".lock", timeout=lock_timeout_seconds)

    def append(self, record: EnquiryRecord) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        try:
            with self._lock:
                workbook = self._open_or_create()
                sheet = workbook[SHEET_NAME]
                sheet.append(record.as_row())
                self._format_row(sheet, sheet.max_row)
                self._atomic_save(workbook)
        except Timeout as exc:
            raise ExcelStoreError("Timed out waiting for the spreadsheet lock") from exc
        except PermissionError as exc:
            # Most often: the file is open in Excel on Windows, which blocks writes.
            raise ExcelStoreError(f"No permission to write {self.path}; is it open in Excel?") from exc
        except OSError as exc:
            raise ExcelStoreError(f"Could not write {self.path}: {exc}") from exc

        logger.info("Stored enquiry %s in %s (row %d)", record.reference, self.path.name, sheet.max_row)

    def _open_or_create(self) -> Workbook:
        if self.path.exists():
            workbook = load_workbook(self.path)
            if SHEET_NAME not in workbook.sheetnames:
                self._init_sheet(workbook.create_sheet(SHEET_NAME))
            return workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        self._init_sheet(sheet)
        logger.info("Created new enquiries workbook at %s", self.path)
        return workbook

    @staticmethod
    def _init_sheet(sheet) -> None:  # type: ignore[no-untyped-def]
        sheet.append([column.header for column in COLUMNS])
        header_fill = PatternFill("solid", fgColor="0B2545")
        for index, column in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            sheet.column_dimensions[get_column_letter(index)].width = column.width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    @staticmethod
    def _format_row(sheet, row_number: int) -> None:  # type: ignore[no-untyped-def]
        for cell in sheet[row_number]:
            if isinstance(cell.value, str):
                # Plain text cell: never interpreted as a formula.
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif isinstance(cell.value, int):
                # Real number, so the team can sort and filter budgets.
                cell.number_format = "#,##0"
                cell.alignment = Alignment(vertical="top")

    def _atomic_save(self, workbook: Workbook) -> None:
        fd, tmp_name = tempfile.mkstemp(dir=self.path.parent, prefix=".enquiries-", suffix=".xlsx")
        os.close(fd)
        try:
            workbook.save(tmp_name)
            os.replace(tmp_name, self.path)
        finally:
            if os.path.exists(tmp_name):
                os.remove(tmp_name)
